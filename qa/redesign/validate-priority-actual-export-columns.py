"""Independently parse bytes generated from the actual UI export column declarations."""
from __future__ import annotations
from decimal import Decimal
import json
from pathlib import Path
import warnings
from zipfile import ZipFile
import openpyxl

project = Path(__file__).resolve().parents[2]
manifest = json.loads(Path('/private/tmp/priority-actual-ui-exports/manifest.json').read_text())
checks = []

def check(name, condition, details=None):
    checks.append({'name': name, 'pass': bool(condition), **({'details': details} if details else {})})

def close(a, b, tolerance=Decimal('0.000001')):
    return abs(Decimal(str(a))-Decimal(str(b))) <= tolerance

loaded = {}
for variant in manifest['variants']:
    with warnings.catch_warnings(record=True) as recorded:
        workbook = openpyxl.load_workbook(variant['path'], data_only=False)
    check(f"{variant['name']}: valid Data/Scope/Definitions and warning-free parse", workbook.sheetnames == ['Data', 'Scope', 'Definitions'] and not recorded)
    with ZipFile(variant['path']) as package:
        check(f"{variant['name']}: package CRC", package.testzip() is None)
    sheet = workbook['Data']
    headers = [cell.value for cell in sheet[1]]
    rows = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=2, values_only=True)]
    loaded[variant['name']] = (workbook, headers, rows)
    check(f"{variant['name']}: actual UI row identity/count", len(rows) == variant['rowCount'] and [row[headers[0]] for row in rows] == [row[headers[0]] for row in variant['expectedRows']])
    numeric_ok = True
    for row_number, expected in enumerate(variant['expectedRows'], 2):
        for column_number, header in enumerate(headers, 1):
            value = expected[header]
            actual = sheet.cell(row_number, column_number)
            if isinstance(value, (int, float)):
                numeric_ok = numeric_ok and actual.data_type == 'n' and close(actual.value, value)
            else:
                numeric_ok = numeric_ok and actual.value == value
    check(f"{variant['name']}: every actual UI value is preserved with numeric cells", numeric_ok)
    scope_rows = dict(workbook['Scope'].iter_rows(min_row=2, values_only=True))
    check(f"{variant['name']}: click-time SS26 scope/source metadata", scope_rows.get('season') == 'SS26' and scope_rows.get('attribute') == 'brand' and scope_rows.get('Exported rows') == len(rows) and scope_rows.get('Source') == 'SS26 coherent flooring demo')
    if variant['scope'] == 'bamboo':
        scope = json.loads(scope_rows['hierarchyScope'])
        check(f"{variant['name']}: Bamboo recipe is explicit", scope[0]['selectedValues'] == ['Bamboo'])
    definition_rows = list(workbook['Definitions'].iter_rows(min_row=2, values_only=True))
    definition_text = '\n'.join(str(cell or '') for row in definition_rows for cell in row)
    check(f"{variant['name']}: period, original baseline and area caveats survive", all(term in definition_text for term in ['4 completed', '13 completed', '52 completed', '8.8M', 'not a recovered client PLR', 'not selling-unit AUR', 'definition TBC']))
    check(f"{variant['name']}: actual caller filename uses SS26", 'SS26' in variant['fileName']['view'] and '74' not in variant['fileName']['view'])
    if variant['rowSelection'] == 'filtered':
        check(f"{variant['name']}: filtered rows and current grid view metadata", len(rows) == 1 and rows[0][headers[0]] == 'Woodland Reserve' and 'greaterThan' in scope_rows.get('View: filters', '') and 'desc' in scope_rows.get('View: sort', ''))

full_book, full_headers, full_rows = loaded['breakdown-full-all']
totals = manifest['expectedFull']
metric_headers = {
    'Sales': 'sales', 'LY sales': 'lyRevenue', 'Plan': 'planRevenue',
    'PLR baseline · definition TBC': 'plrBaseline',
    'R4 sales · definition TBC': 'r4Sales', 'R13 sales · definition TBC': 'r13Sales',
    'R52 sales · definition TBC': 'r52Sales',
}
check('Actual lower columns export the exact $8.4M full-scope sales', sum(Decimal(str(row['Sales'])) for row in full_rows) == Decimal('8400000'))
check('Actual lower columns export exact LY/Plan/PLR/R4/R13/R52 cents', all(close(sum(Decimal(str(row[header])) for row in full_rows), totals[key]) for header, key in metric_headers.items()))
check('Actual full lower export contains the five canonical brands', {row['Sales Mix'] for row in full_rows} == {'Woodland Reserve', 'Westridge', 'American Spirit', 'Lifescapes', 'Bruce'})
bamboo_book, bamboo_headers, bamboo_rows = loaded['breakdown-bamboo-all']
check('Actual Bamboo columns reconcile the selected scope and exact $1,049,808.28 sales', sum(Decimal(str(row['Sales'])) for row in bamboo_rows) == Decimal('1049808.28') and all(close(sum(Decimal(str(row[header])) for row in bamboo_rows), manifest['expectedBamboo'][key]) for header, key in metric_headers.items()))
check('Actual upper/lower shares agree without converting whole-percent values to fractions', all(close(next(row for row in loaded['summary-full-all'][2] if row['Name'] == lower['Sales Mix'])['Sales share %'], lower['Sales Share']) for lower in full_rows) and close(next(row for row in full_rows if row['Sales Mix'] == 'Woodland Reserve')['Sales Share'], 27.4))
check('Percentage number format displays stored whole-percent values correctly', all('"%"' in full_book['Data'].cell(2, full_headers.index(header)+1).number_format for header in ['Growth', 'Sales Share', 'SKU Share', 'GM %', 'Returns']))
check('Sales-per-area exports the weighted brand ratios rather than adding ratios', all(close(row['Sales per sq ft · definition TBC'], expected['Sales per sq ft · definition TBC']) for row, expected in zip(full_rows, next(v for v in manifest['variants'] if v['name']=='breakdown-full-all')['expectedRows'])))

report = {
    'kind': manifest['kind'],
    'status': 'passed' if all(check['pass'] for check in checks) else 'failed',
    'passed': sum(check['pass'] for check in checks), 'checkCount': len(checks),
    'actual_browser_download_verified': False,
    'source_extraction': manifest['extraction'],
    'source_snapshots': manifest['sources'],
    'variants': [{key: value for key, value in variant.items() if key not in ['expectedRows']} for variant in manifest['variants']],
    'checks': checks,
    'limits': [
        'Executes the actual current UI column functions and metadata arrows with synthetic full/Bamboo rows; React and AG Grid are not mounted.',
        'Filtered-row variants verify writer fidelity to caller inputs; actual getDisplayedRows traversal, browser click and saved download still require browser acceptance.',
        'The upper summary export intentionally contains name, SKU count, sales share, LY change and productivity; C4 period and baseline columns belong to the lower breakdown export.',
    ],
}
output = project / 'qa/redesign/priority-actual-export-columns-checks.json'
output.write_text(json.dumps(report, indent=2, ensure_ascii=False)+'\n')
print(json.dumps({'status': report['status'], 'passed': report['passed'], 'checkCount': report['checkCount'], 'report': str(output)}, indent=2))
raise SystemExit(0 if report['status'] == 'passed' else 1)
