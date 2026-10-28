"""Independent XLSX parse of generated model controls; not native download proof."""
from pathlib import Path
import json, re
from decimal import Decimal
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
DIR=ROOT/'qa/redesign/mix-export-model-controls'
results=[]

def records(ws):
    values=list(ws.values)
    return [dict(zip(values[0],row)) for row in values[1:]] if values else []

def amount(rows,key):
    values=[row[key] for row in rows if row.get(key) is not None]
    assert all(isinstance(value,(int,float)) for value in values),(key,'Numeric metric became text')
    return sum((Decimal(str(value)) for value in values),Decimal(0))

for file in sorted(DIR.glob('*.xlsx')):
    wb=load_workbook(file,data_only=False)
    assert wb.sheetnames==['Data','Scope','Definitions','Contributing pairs']
    rows=records(wb['Data']);leaves=records(wb['Contributing pairs'])
    scope={row['Scope field']:row['Value'] for row in records(wb['Scope'])}
    defs=records(wb['Definitions'])
    denominator=json.loads(scope['View: shareDenominator'])
    assert scope['Source version']=='flooring-ledger-ss26-v2'
    assert scope['window']=='SEASON' and scope['stockAsOf']=='2026-07-31'
    assert scope['rangeAsOf']=='2026-07-31' and scope['uom']=='SELLING_UNIT_TBC'
    assert scope['Contributing pair rows']==len(leaves)
    assert scope['Exported rows']==len(rows)
    assert len({row['pairId'] for row in leaves})==len(leaves)
    assert all(row['pairId']==f"{row['itemId']}:{row['storeId']}" for row in leaves)
    assert not any(cell.data_type=='f' for ws in wb for row in ws for cell in row),'No formulas should be introduced from text.'
    if file.stem=='all-brands':
        assert len(rows)==5 and len(leaves)==3240 and len(defs)==87
        assert amount(rows,'sales')==amount(leaves,'sales')==Decimal('8400000')
        assert amount(rows,'quantity')==amount(leaves,'quantity')==Decimal('182400')
    elif file.stem in ['woodland','restricted-woodland']:
        assert len(rows)==1 and rows[0]['id']=='Woodland Reserve' and rows[0]['skuCount']==19
        assert len(leaves)==855 and len({row['itemId'] for row in leaves})==19
        assert len({row['storeId'] for row in leaves})==45
        assert amount(rows,'sales')==amount(leaves,'sales')==Decimal('2301600')
        assert amount(rows,'quantity')==amount(leaves,'quantity')==Decimal('44861')
        assert rows[0]['salesShare']==27.4 and denominator['sales']['value']==8400000
        assert denominator['pairCount']==3240
    elif file.stem=='exact-diagonal':
        assert len(leaves)==2
        assert {row['pairId'] for row in leaves}=={'1000037:ST1001','1000444:ST1002'}
        assert amount(leaves,'sales')==Decimal('52285.80')
    elif file.stem=='no-matching-groups':
        assert not rows and not leaves and denominator['pairCount']==3240
    if file.stem=='restricted-woodland':
        assert scope['access'] in ['restricted-margin','[Restricted margin evidence]']
        assert 'fields withheld' in scope['Access policy']
        denied=re.compile(r'(^|[ _.-])(gm|cogs|margin|gross.?profit)([ _.$%.-]|$)|gmDollars|gmRate|plannedGmRate|gmGap|recognizedCogs|grossMargin',re.I)
        for sheet in ['Data','Contributing pairs']:
            assert not any(denied.search(str(header)) for header in next(wb[sheet].values))
        assert not any(denied.search(str(row['Metric'])) for row in defs)
        assert len(defs)<87
    else:
        assert len(defs)==87
    results.append({'file':file.name,'dataRows':len(rows),'leafRows':len(leaves),'definitions':len(defs),'passed':True})

assert len(results)==5
report={'status':'passed','workbooks':5,'method':'Independent openpyxl read of generated controls from the actual MixView export adapter. No browser download is claimed.','results':results}
(ROOT/'qa/redesign/mix-export-xlsx-checks.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps(report))
