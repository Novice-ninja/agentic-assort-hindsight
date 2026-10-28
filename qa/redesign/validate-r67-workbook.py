import json
import warnings
import zipfile
from pathlib import Path

import openpyxl

base=Path(__file__).resolve().parents[2]
folder=base/'qa/redesign/r67-export-checks'
catalogue=json.loads((base/'src/modules/hindsight/redesign/data/catalogue.json').read_text())
checks=[]
def check(name,value):
    checks.append({'name':name,'pass':bool(value)})
    assert value,name

for access in ['full','restricted-margin']:
    path=folder/f'{access}.xlsx'
    with warnings.catch_warnings(record=True) as caught:
        book=openpyxl.load_workbook(path,data_only=True)
    check(f'{access}: warning-free load',not caught)
    check(f'{access}: ZIP CRC',zipfile.ZipFile(path).testzip() is None)
    check(f'{access}: four worksheets',book.sheetnames==['Data','Scope','Definitions','Contributing pairs'])
    data=list(book['Data'].values)
    row=dict(zip(data[0],data[1]))
    check(f'{access}: text identifier',row['id']=='1000037')
    check(f'{access}: exact numeric sales',row['sales']==1234.56)
    check(f'{access}: missing distinct',row['nullMetric'] is None and row['nullMetric state']=='missing')
    check(f'{access}: XML controls cleaned',row['note']=='XML control & <characters>')
    scope=dict(list(book['Scope'].values)[1:])
    check(f'{access}: metric version',scope['Metric version']=='proposed-base-v2')
    defs=list(book['Definitions'].values)
    records=[dict(zip(defs[0],row)) for row in defs[1:]]
    for metric_id in ['K01','K04'] if access=='full' else ['K01']:
        original=next(d for d in catalogue if d['id']==metric_id)
        record=next(d for d in records if d['Metric ID']==metric_id)
        check(f'{access}: {metric_id} exact evidence status',record['Status']==original['evidence_status'])
        check(f'{access}: {metric_id} aliases retained',all(alias.strip() in record['Source aliases'] for alias in original['source_labels'].split(';')))
        check(f'{access}: {metric_id} aggregation guards',record['Aggregation guards']==original['aggregation_guards'])
        check(f'{access}: {metric_id} definition version',record['Definition version']=='proposed-base-v2')
    contributor_rows=list(book['Contributing pairs'].values)
    contributor=dict(zip(contributor_rows[0],contributor_rows[1]))
    check(f'{access}: valid stock zero',contributor['sellableOh']==0 and contributor['sellableOh state']=='ok')
    if access=='full':
        check('full: 87 metric definitions',len(records)==87)
        check('full: numeric GM retained',row['gmRate']==30.0001)
    else:
        check('restricted: GM columns removed','gmRate' not in row and 'gmDollars' not in contributor)
        check('restricted: recursive row fields removed','789' not in row['nested'])
        check('restricted: recursive view fields removed','42' not in str(scope.get('View: grossMargin','')) and '51' not in scope['View: nested'])

book=openpyxl.load_workbook(folder/'large-scope.xlsx',data_only=True)
scope=dict(list(book['Scope'].values)[1:])
expected=json.loads((folder/'large-scope-expected.json').read_text())
for field in ['pairIds','View: exactPairs']:
    parts=[value for key,value in scope.items() if key.startswith(field+' (part ')]
    check(f'{field}: exact metadata restored',json.loads(''.join(parts))==expected)
check('large scope: every cell within Excel character limit',all(not isinstance(cell,str) or len(cell)<=32767 for sheet in book for row in sheet.values for cell in row))
(folder/'validation.json').write_text(json.dumps({'checks':checks,'passed':len(checks),'browser':False,'note':'Actual adapter bytes parsed independently with openpyxl; no UI export click claim.'},indent=2)+'\n')
print(f'{len(checks)} independent workbook checks passed')
