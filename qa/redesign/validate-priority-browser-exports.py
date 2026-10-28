"""Read-only independent audit of actual XLSX files saved from native browser downloads."""
from datetime import datetime
from decimal import Decimal
from hashlib import sha256
import json
from pathlib import Path
import warnings
from zipfile import ZipFile
import openpyxl

project = Path(__file__).resolve().parents[2]
expected = json.loads((project/'qa/redesign/priority-browser-export-expectations.json').read_text())
checks, files, books = [], [], {}
def check(name, condition, details=None):
    checks.append({'name':name,'pass':bool(condition),**({'details':details} if details is not None else {})})
def near(a,b,tolerance=Decimal('0.000001')):
    return isinstance(a,(int,float,Decimal)) and abs(Decimal(str(a))-Decimal(str(b))) <= tolerance
def total(rows,key):
    return sum((Decimal(str(row[key])) for row in rows),Decimal(0))

for relative, want in expected.items():
    path=project/relative
    original_hash=sha256(path.read_bytes()).hexdigest()
    with warnings.catch_warnings(record=True) as recorded:
        workbook=openpyxl.load_workbook(path,data_only=False)
    check(path.name+': valid three-sheet XLSX, no parser warnings',workbook.sheetnames==['Data','Scope','Definitions'] and not recorded)
    with ZipFile(path) as package:
        check(path.name+': valid package CRC',package.testzip() is None)
    sheet=workbook['Data'];headers=[cell.value for cell in sheet[1]]
    rows=[dict(zip(headers,values)) for values in sheet.iter_rows(min_row=2,values_only=True)]
    scope=dict(workbook['Scope'].iter_rows(min_row=2,values_only=True))
    definitions=list(workbook['Definitions'].iter_rows(min_row=2,values_only=True))
    identity='Style ID' if path.name.startswith('by-item-') else 'Name' if 'summary' in path.name else 'Sales Mix'
    expected_rows={str(row[identity]):row for row in want['rows']}
    actual_rows={str(row[identity]):row for row in rows}
    check(path.name+': exact canonical row identities/count, no duplicate rollups',len(rows)==len(expected_rows)==len(actual_rows) and set(actual_rows)==set(expected_rows))
    check(path.name+': exact current UI export headers',headers==list(want['rows'][0]))
    mismatches=[];numeric_cells=0;types_ok=True
    for row_number,row in enumerate(rows,2):
        reference=expected_rows.get(str(row[identity]),{})
        for column_number,header in enumerate(headers,1):
            value=reference.get(header)
            cell=sheet.cell(row_number,column_number)
            if isinstance(value,(int,float)):
                numeric_cells+=1
                types_ok=types_ok and cell.data_type=='n'
                if not near(cell.value,value):mismatches.append({'row':row_number,'header':header,'actual':cell.value,'expected':value})
            elif cell.value!=value:
                mismatches.append({'row':row_number,'header':header,'actual':cell.value,'expected':value})
    check(path.name+': every exported value agrees with real API/actual column adapter',not mismatches,mismatches[:8] if mismatches else {'compared_rows':len(rows),'compared_columns':len(headers)})
    check(path.name+': raw numeric cells remain typed numbers',types_ok and numeric_cells>0,{'numeric_cells':numeric_cells})
    check(path.name+': no formula/error cells introduced',not any(cell.data_type in ('f','e') for row in sheet for cell in row))
    check(path.name+': exported count matches Data',scope.get('Exported rows')==len(rows))
    check(path.name+': SS26, source/version and explicit recipe metadata',scope.get('season')=='SS26' and scope.get('planCode')==want['scope']['planCode'] and scope.get('Source')=='SS26 coherent flooring demo' and bool(scope.get('Source version')) and json.loads(scope['hierarchyScope'])==want['scope']['hierarchyScope'])
    check(path.name+': generated filename names SS26, not the reference code', 'SS26' in scope.get('Export','') and '_74_' not in scope.get('Export',''))
    try:datetime.fromisoformat(scope['Exported at'].replace('Z','+00:00'));timestamp_valid=True
    except (KeyError,ValueError,TypeError):timestamp_valid=False
    check(path.name+': export timestamp is a valid ISO instant',timestamp_valid)
    text='\n'.join(str(value or '') for row in definitions for value in row)
    check(path.name+': definitions disclose periods, area denominator, original baseline and synthetic status',all(term in text for term in ['4 completed','13 completed','52 completed','selling area','not selling-unit AUR','not a recovered client PLR','Not client actuals','Synthetic · definition TBC']))
    if identity=='Style ID':
        check(path.name+': stable IDs remain text and match all scoped canonical products',set(actual_rows)==set(want['canonicalIds']) and all(sheet.cell(i,1).data_type in ('s','inlineStr') for i in range(2,sheet.max_row+1)))
        check(path.name+': loaded/exported/scoped counts disclose exact export population',all(scope.get('View: '+field)==len(rows) for field in ['loadedRows','exportedRows','scopedItemsBeforeLocalTableFilters']) and 'loaded filtered and sorted product leaf rows' in scope.get('View: exportPopulation','') and scope.get('View: moreRowsAvailable')=='false')
        check(path.name+': plan and reference season are distinct metadata fields',scope.get('planSeasonCode')==78 and scope.get('referenceSeasonCode')==74)
        metric_headers={'Sls$':'sales','SlsU':'units','LY sales $':'lyRevenue','Plan sales $ · synthetic':'planRevenue','PLR baseline $ · definition TBC':'plrBaseline','R4 sales $ · definition TBC':'r4Sales','R13 sales $ · definition TBC':'r13Sales','R52 sales $ · definition TBC':'r52Sales','GM$':'grossMargin','Buy Units':'buyUnits','Buy $':'buyDollars'}
    elif identity=='Sales Mix':
        metric_headers={'Sales':'sales','Selling SKUs':'itemCount','LY sales':'lyRevenue','Plan':'planRevenue','PLR baseline · definition TBC':'plrBaseline','R4 sales · definition TBC':'r4Sales','R13 sales · definition TBC':'r13Sales','R52 sales · definition TBC':'r52Sales'}
        current_filters=json.loads(scope.get('View: filters','null'))
        current_sort=json.loads(scope.get('View: sort','null'))
        check(path.name+': all current grid filter/sort fields are machine readable',isinstance(current_filters,dict) and isinstance(current_sort,list))
        if 'sorted' in path.name:
            check(path.name+': actual Sales ascending metadata is captured',current_filters=={} and current_sort==[{'column':'sales','direction':'asc','order':0}])
            check(path.name+': actual row order follows visible Sales ascending sort',[row['Sales Mix'] for row in rows]==['Bruce','Lifescapes','American Spirit','Westridge','Woodland Reserve'] and [row['Sales'] for row in rows]==sorted(row['Sales'] for row in rows))
        check(path.name+': sales and SKU shares each sum to100%',near(total(rows,'Sales Share'),100) and near(total(rows,'SKU Share'),100))
        weighted_gm=sum(Decimal(str(row['Sales']))*Decimal(str(row['GM %'])) for row in rows)/total(rows,'Sales')
        check(path.name+': sales-weighted GM matches canonical ratio',near(weighted_gm,want['totals']['gmPct']))
    else:
        metric_headers={'Total SKUs':'itemCount'}
        check(path.name+': all five brands and whole-percent shares are preserved',len(rows)==5 and near(total(rows,'Sales share %'),100) and near(actual_rows['Woodland Reserve']['Sales share %'],27.4))
    check(path.name+': all additive monetary/quantity/period totals reconcile',all(near(total(rows,header),want['totals'][key]) for header,key in metric_headers.items()),{header:float(total(rows,header)) for header in metric_headers})
    if identity!='Style ID':
        percent_headers=['Sales share %','Vs LY'] if identity=='Name' else ['Growth','Sales Share','SKU Share','GM %','Returns']
        check(path.name+': percent number format preserves stored whole-percent semantics',all('"%"' in sheet.cell(2,headers.index(header)+1).number_format for header in percent_headers))
    check(path.name+': workbook bytes were not changed by audit',sha256(path.read_bytes()).hexdigest()==original_hash)
    files.append({'path':relative,'sha256':original_hash,'row_count':len(rows),'column_count':len(headers),'numeric_cells_checked':numeric_cells,'canonical_comparison':want['source'],'scope_recipe':want['scope']['hierarchyScope'],'totals':{header:float(total(rows,header)) for header in metric_headers}})
    books[path.name]={'rows':rows,'scope':scope,'totals':want['totals']}

full=books['by-item-full.xlsx'];bamboo=books['by-item-bamboo.xlsx']
check('Full By Item exact headline controls',len(full['rows'])==72 and total(full['rows'],'Sls$')==Decimal('8400000') and total(full['rows'],'SlsU')==Decimal('182400'))
check('Bamboo By Item exact headline controls',len(bamboo['rows'])==12 and total(bamboo['rows'],'Sls$')==Decimal('1049808.28') and total(bamboo['rows'],'SlsU')==Decimal('23836'))
check('Bamboo contains only its12 products and excludes Northport',all(row['Department']=='Bamboo' for row in bamboo['rows']) and '1000037' not in {row['Style ID'] for row in bamboo['rows']} and '1000037' in {row['Style ID'] for row in full['rows']} )
for scope_name in ['full','bamboo']:
    by_item=books['by-item-'+scope_name+'.xlsx'];mix=books['sales-mix-breakdown-'+scope_name+'.xlsx']
    check(scope_name+': By Item and Sales Mix exact same scope recipe',json.loads(by_item['scope']['hierarchyScope'])==json.loads(mix['scope']['hierarchyScope']))
    check(scope_name+': By Item and Sales Mix reconcile sales/LY/Plan/PLR/R4/R13/R52 across actual browser files',all(near(total(by_item['rows'],a),total(mix['rows'],b)) for a,b in [('Sls$','Sales'),('LY sales $','LY sales'),('Plan sales $ · synthetic','Plan'),('PLR baseline $ · definition TBC','PLR baseline · definition TBC'),('R4 sales $ · definition TBC','R4 sales · definition TBC'),('R13 sales $ · definition TBC','R13 sales · definition TBC'),('R52 sales $ · definition TBC','R52 sales · definition TBC')]))
summary=books['sales-mix-summary-full.xlsx'];lower=books['sales-mix-breakdown-full.xlsx']
check('Full summary and lower brand shares/SKU counts/productivity agree',all(any(other['Sales Mix']==row['Name'] and near(other['Sales Share'],row['Sales share %']) and near(other['Selling SKUs'],row['Total SKUs']) and near(other['Productivity'],row['Productivity']) for other in lower['rows']) for row in summary['rows']))

report={'kind':'Independent read-only validation of actual XLSX browser downloads saved through native Save dialogs','status':'passed' if all(c['pass'] for c in checks) else 'failed','passed':sum(c['pass'] for c in checks),'checkCount':len(checks),'browser_files_verified':len(files),'files':files,'checks':checks,'limits':['Five files capture unfiltered local grid states under full/Bamboo global scopes; a sixth captures actual full Brand Sales ascending order and nonempty Scope sort metadata. A nonempty local column filter is not demonstrated; root reports that this particular name-column menu has no filter.','Browser save/click/sort provenance is supplied by root; this agent independently reads the resulting file bytes and real API/column-adapter expectations without operating Chrome.','By Item leaf rates may be rounded by its API contract. Aggregate ratios are recomputed from additive evidence; individual rates or sales-per-area values are never summed.','These are synthetic local prototype exports; production fiscal periods, quantity UOM, selling area and PLR definitions remain explicitly TBC.']}
out=project/'qa/redesign/priority-browser-export-validation.json';out.write_text(json.dumps(report,indent=2,ensure_ascii=False)+'\n')
print(json.dumps({'status':report['status'],'passed':report['passed'],'checks':len(checks),'files':len(files),'failed_checks':[c['name'] for c in checks if not c['pass']]},indent=2))
raise SystemExit(0 if report['status']=='passed' else 1)
